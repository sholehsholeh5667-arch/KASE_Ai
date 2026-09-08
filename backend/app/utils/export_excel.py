"""
Export Excel
============

Utility untuk membuat laporan Excel.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """
    Export laporan ke Excel.
    """

    def __init__(self):

        self.header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        self.header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        self.center = Alignment(
            horizontal="center"
        )

    # ==========================================
    # Export
    # ==========================================

    def export(
        self,
        title: str,
        headers: list,
        rows: list
    ):
        """
        Membuat file Excel.
        """

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Laporan"

        # ==========================================
        # Judul
        # ==========================================

        sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(headers)
        )

        cell = sheet.cell(
            row=1,
            column=1
        )

        cell.value = title

        cell.font = Font(
            bold=True,
            size=14
        )

        cell.alignment = self.center

        # ==========================================
        # Header
        # ==========================================

        for index, header in enumerate(
            headers,
            start=1
        ):

            c = sheet.cell(
                row=3,
                column=index
            )

            c.value = header

            c.font = self.header_font

            c.fill = self.header_fill

            c.alignment = self.center

        # ==========================================
        # Data
        # ==========================================

        row_index = 4

        for row in rows:

            for col_index, value in enumerate(
                row,
                start=1
            ):

                sheet.cell(
                    row=row_index,
                    column=col_index
                ).value = value

            row_index += 1

        # ==========================================
        # Auto Width
        # ==========================================

        # Jangan menggunakan:
        #
        # for column_cells in sheet.columns:
        #     column = column_cells[0].column_letter
        #
        # karena baris judul menggunakan merged cells.
        #
        # Gunakan nomor kolom secara langsung agar aman
        # terhadap MergedCell.

        for col_index in range(
            1,
            sheet.max_column + 1
        ):

            column = get_column_letter(
                col_index
            )

            length = 0

            for row_index in range(
                1,
                sheet.max_row + 1
            ):

                current_cell = sheet.cell(
                    row=row_index,
                    column=col_index
                )

                if current_cell.value is not None:

                    length = max(
                        length,
                        len(str(current_cell.value))
                    )

            sheet.column_dimensions[
                column
            ].width = length + 4

        # ==========================================
        # Simpan ke Memory
        # ==========================================

        buffer = BytesIO()

        workbook.save(buffer)

        excel = buffer.getvalue()

        buffer.close()

        return excel